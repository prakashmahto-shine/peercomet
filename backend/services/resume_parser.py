"""
Resume text extraction (PDF / DOCX / TXT) and rule-based field parsing.

Parsing is best-effort and offline (no external AI service). Anything not found is left
out rather than guessed: the user reviews the pre-filled profile before saving.
"""
import io
import re
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

MAX_TEXT_CHARS = 200_000


class ResumeTextError(Exception):
    """The file couldn't be read as a resume of the given type."""


# ----------------------------------------------------------------------------- extraction
def extract_text(data: bytes, extension: str) -> str:
    ext = extension.lower().lstrip(".")
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                reader.decrypt("")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif ext == "docx":
            import docx
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        elif ext == "txt":
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("latin-1")
        else:
            raise ResumeTextError(f"Unsupported file type: .{ext}")
    except ResumeTextError:
        raise
    except Exception as e:
        raise ResumeTextError(f"Could not read the {ext.upper()} file") from e
    return text[:MAX_TEXT_CHARS]


# ----------------------------------------------------------------------------- vocabularies
SKILLS = [
    # Languages
    "Python", "Java", "JavaScript", "TypeScript", "C", "C++", "C#", "Go", "Golang", "Rust", "Kotlin", "Swift",
    "PHP", "Ruby", "Scala", "R", "MATLAB", "Dart", "Perl", "Shell Scripting", "Bash", "SQL", "PL/SQL",
    # Frontend
    "React", "React.js", "React Native", "Next.js", "Angular", "Vue.js", "Svelte", "Redux", "HTML", "HTML5",
    "CSS", "CSS3", "Sass", "Tailwind CSS", "Bootstrap", "jQuery", "Webpack", "Vite", "Micro-Frontends",
    "Module Federation", "Web Performance", "Responsive Design", "Flutter",
    # Backend
    "Node.js", "Express.js", "NestJS", "Django", "Flask", "FastAPI", "Spring", "Spring Boot", "Hibernate",
    ".NET", "ASP.NET", "Laravel", "Ruby on Rails", "GraphQL", "REST APIs", "gRPC", "Microservices",
    "Kafka", "RabbitMQ", "Redis", "Elasticsearch", "Solr", "Lucene",
    # Data / databases
    "MySQL", "PostgreSQL", "MongoDB", "Cassandra", "DynamoDB", "Oracle", "SQL Server", "SQLite", "Snowflake",
    "BigQuery", "Hadoop", "Spark", "PySpark", "Airflow", "dbt", "ETL", "Data Warehousing", "Tableau",
    "Power BI", "Excel", "Data Analysis", "Data Visualization", "Statistics",
    # AI / ML
    "Machine Learning", "Deep Learning", "NLP", "Computer Vision", "TensorFlow", "PyTorch", "Keras",
    "scikit-learn", "Pandas", "NumPy", "LLMs", "Generative AI", "RAG", "LangChain", "MLOps",
    # Cloud / DevOps
    "AWS", "Azure", "GCP", "Google Cloud", "Docker", "Kubernetes", "Terraform", "Ansible", "Jenkins",
    "CI/CD", "GitHub Actions", "Git", "Linux", "Nginx", "Prometheus", "Grafana", "Serverless",
    # Security
    "Cybersecurity", "Cloud Security", "DevSecOps", "Penetration Testing", "SIEM", "Splunk", "IAM",
    "Network Security", "OWASP",
    # Semiconductor
    "Verilog", "SystemVerilog", "VHDL", "UVM", "RTL Design", "ASIC", "FPGA", "STA", "Physical Design",
    "Design Verification",
    # Testing
    "Selenium", "Cypress", "Jest", "JUnit", "Pytest", "Automation Testing", "Manual Testing", "QA",
    # Architecture / practices
    "System Design", "Distributed Systems", "Data Structures", "Algorithms", "Agile", "Scrum", "JIRA",
    "Object-Oriented Programming", "Design Patterns",
    # Product / business
    "Product Management", "Product Strategy", "Roadmapping", "A/B Testing", "Product Analytics",
    "Stakeholder Management", "Project Management", "Program Management", "Business Analysis",
    "Digital Marketing", "SEO", "SEM", "Content Marketing", "Social Media Marketing", "Growth Marketing",
    "Sales", "B2B Sales", "SaaS Sales", "Account Management", "Business Development", "CRM", "Salesforce",
    "HubSpot", "Customer Success", "Recruitment", "Talent Acquisition", "HR", "Finance", "Accounting",
    "Financial Analysis", "Figma", "UI/UX", "UX Research", "Adobe Photoshop", "Communication", "Leadership",
]

# Ordered from highest to lowest
DEGREES: List[Tuple[str, str]] = [
    (r"\bph\.?\s?d\b|\bdoctorate\b", "PhD"),
    (r"\bm\.?\s?tech\b", "M.Tech"),
    (r"\bm\.?\s?e\.?(?=\s|,|$)", "M.E."),
    (r"\bmba\b|\bpgdm\b", "MBA"),
    (r"\bmca\b", "MCA"),
    (r"\bm\.?\s?sc\b", "M.Sc"),
    (r"\bm\.?\s?com\b", "M.Com"),
    (r"\bm\.?\s?a\.?(?=\s|,|$)", "M.A."),
    (r"\bmaster'?s?\b|\bmasters? of\b", "Master's"),
    (r"\bb\.?\s?tech\b", "B.Tech"),
    (r"\bb\.?\s?e\.?(?=\s|,|$)", "B.E."),
    (r"\bbca\b", "BCA"),
    (r"\bb\.?\s?sc\b", "B.Sc"),
    (r"\bb\.?\s?com\b", "B.Com"),
    (r"\bbba\b", "BBA"),
    (r"\bb\.?\s?a\.?(?=\s|,|$)", "B.A."),
    (r"\bbachelor'?s?\b|\bbachelors? of\b", "Bachelor's"),
    (r"\bdiploma\b", "Diploma"),
]

CITIES = [
    "Bengaluru", "Bangalore", "Mumbai", "Navi Mumbai", "Delhi", "New Delhi", "Gurugram", "Gurgaon", "Noida",
    "Greater Noida", "Hyderabad", "Chennai", "Pune", "Kolkata", "Ahmedabad", "Jaipur", "Chandigarh",
    "Kochi", "Coimbatore", "Indore", "Lucknow", "Nagpur", "Bhubaneswar", "Thiruvananthapuram", "Mysuru",
    "Vadodara", "Surat", "Bhopal", "Visakhapatnam", "Mohali", "Faridabad", "Ghaziabad", "Remote",
]

TITLE_WORDS = [
    "engineer", "developer", "manager", "analyst", "designer", "lead", "architect", "consultant", "intern",
    "scientist", "specialist", "executive", "officer", "head", "director", "administrator", "associate",
    "programmer", "tester", "devops", "sde", "president", "founder", "recruiter", "accountant",
]

SECTION_HEADINGS = [
    "summary", "professional summary", "profile", "profile summary", "objective", "career objective",
    "about", "about me", "experience", "work experience", "professional experience", "employment history",
    "employment", "education", "academic", "academics", "qualifications", "skills", "technical skills",
    "key skills", "core competencies", "projects", "certifications", "achievements", "awards",
    "languages", "interests", "hobbies", "personal details", "personal information", "declaration",
    "contact", "internships",
]

MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}


# ----------------------------------------------------------------------------- helpers
def _lines(text: str) -> List[str]:
    return [re.sub(r"\s+", " ", l).strip() for l in text.splitlines() if l.strip()]


def _heading_of(line: str) -> Optional[str]:
    clean = re.sub(r"[^a-z ]", "", line.lower()).strip()
    return clean if clean in SECTION_HEADINGS and len(line) <= 40 else None


def _section(lines: List[str], names: List[str]) -> List[str]:
    """Lines under the first heading in `names`, up to the next heading."""
    for i, line in enumerate(lines):
        if _heading_of(line) in names:
            out = []
            for nxt in lines[i + 1:]:
                if _heading_of(nxt):
                    break
                out.append(nxt)
            return out
    return []


def _parse_month_year(token: str) -> Optional[Tuple[int, int]]:
    token = token.strip().lower()
    if token in ("present", "current", "now", "till date", "to date", "date", "ongoing"):
        today = date.today()
        return today.year, today.month
    m = re.match(r"([a-z]{3})[a-z]*\.?\s*'?(\d{2,4})$", token)
    if m and m.group(1) in MONTHS:
        year = int(m.group(2))
        year = year + 2000 if year < 100 else year
        return year, MONTHS[m.group(1)]
    m = re.match(r"(\d{1,2})[/\-.](\d{2,4})$", token)
    if m and 1 <= int(m.group(1)) <= 12:
        year = int(m.group(2))
        year = year + 2000 if year < 100 else year
        return year, int(m.group(1))
    m = re.match(r"(\d{4})$", token)
    if m:
        return int(m.group(1)), 1
    return None


_DATE_TOKEN = r"(?:[A-Za-z]{3,9}\.?\s*'?\d{2,4}|\d{1,2}[/\-.]\d{2,4}|\d{4})"
_RANGE_RE = re.compile(
    rf"({_DATE_TOKEN})\s*(?:-|–|—|to|till)\s*({_DATE_TOKEN}|present|current|now|till date|to date|ongoing)",
    re.IGNORECASE,
)


def _experience_from_ranges(lines: List[str]) -> Optional[int]:
    """Total months from date ranges in the experience section (overlaps merged)."""
    section = _section(lines, ["experience", "work experience", "professional experience",
                               "employment history", "employment", "internships"])
    if not section:
        return None
    intervals = []
    for line in section:
        for start_tok, end_tok in _RANGE_RE.findall(line):
            start, end = _parse_month_year(start_tok), _parse_month_year(end_tok)
            if not start or not end:
                continue
            s, e = start[0] * 12 + start[1], end[0] * 12 + end[1]
            if 1970 * 12 <= s <= e <= (date.today().year + 1) * 12:
                intervals.append((s, e))
    if not intervals:
        return None
    intervals.sort()
    total, (cur_s, cur_e) = 0, intervals[0]
    for s, e in intervals[1:]:
        if s <= cur_e:
            cur_e = max(cur_e, e)
        else:
            total += cur_e - cur_s
            cur_s, cur_e = s, e
    total += cur_e - cur_s
    return total if total > 0 else None


def _to_lpa(amount: str, unit: str) -> Optional[float]:
    try:
        value = float(amount.replace(",", ""))
    except ValueError:
        return None
    unit = (unit or "").lower()
    if unit.startswith("cr"):
        value *= 100
    elif unit == "k":
        value = value * 1000 / 100000
    elif not unit and value >= 10000:  # plain rupees, e.g. 1200000
        value = value / 100000
    return round(value, 2) if 0 < value < 1000 else None


# ----------------------------------------------------------------------------- field parsers
def _name(lines: List[str]) -> Tuple[Optional[str], Optional[str]]:
    for line in lines[:6]:
        clean = re.sub(r"^(name\s*[:\-]\s*)", "", line, flags=re.IGNORECASE).strip()
        words = clean.split()
        if not 2 <= len(words) <= 4:
            continue
        if re.search(r"@|\d|resume|curriculum|vitae|\bcv\b|profile|address|phone|email", clean, re.IGNORECASE):
            continue
        if not all(re.fullmatch(r"[A-Za-z][A-Za-z.'\-]*", w) for w in words):
            continue
        if not (clean.istitle() or clean.isupper()):
            continue
        words = [w.capitalize() if w.isupper() else w for w in words]
        return words[0], " ".join(words[1:])
    return None, None


def _experience(text: str, lines: List[str]) -> Tuple[Optional[int], Optional[int]]:
    patterns = [
        r"(\d{1,2}(?:\.\d{1,2})?)\s*\+?\s*(?:years?|yrs?)\s*(?:(?:and\s*)?(\d{1,2})\s*(?:months?|mos?))?\s*(?:of\s+)?(?:[a-z\-]+\s+){0,4}?(?:experience|exp\b)",
        r"(?:total\s+)?experience\s*[:\-]?\s*(\d{1,2}(?:\.\d{1,2})?)\s*\+?\s*(?:years?|yrs?)(?:\s*(\d{1,2})\s*(?:months?|mos?))?",
    ]
    for pattern in patterns:
        m = re.search(pattern, text, re.IGNORECASE)
        if m:
            years_f = float(m.group(1))
            if years_f > 50:
                continue
            years = int(years_f)
            months = int(m.group(2)) if m.group(2) else int(round((years_f - years) * 12))
            return years, min(months, 11)
    months_total = _experience_from_ranges(lines)
    if months_total:
        return months_total // 12, months_total % 12
    return None, None


def _ctc(text: str, kind: str) -> Optional[float]:
    m = re.search(
        rf"{kind}\s*(?:ctc|salary|package)\s*[:\-]?\s*(?:inr|rs\.?|₹)?\s*([\d.,]+)\s*(lpa|lakhs?|lacs?|l\b|crores?|cr|k)?",
        text, re.IGNORECASE,
    )
    return _to_lpa(m.group(1), m.group(2) or "") if m else None


def _notice(text: str) -> Optional[str]:
    m = re.search(r"notice\s*period\s*[:\-]?\s*(immediate(?:ly)?(?:\s*joiner)?|\d{1,3}\s*(?:days?|weeks?|months?))",
                  text, re.IGNORECASE)
    if not m:
        return None
    value = m.group(1).strip()
    return "Immediate" if value.lower().startswith("immediate") else value


def _dob(text: str) -> Optional[str]:
    m = re.search(r"(?:date\s*of\s*birth|d\.?\s?o\.?\s?b\.?)\s*[:\-]?\s*(\d{1,2})(?:st|nd|rd|th)?[/\-. ]([A-Za-z]{3,9}|\d{1,2})[/\-., ]+(\d{4})",
                  text, re.IGNORECASE)
    if not m:
        return None
    day, month_tok, year = int(m.group(1)), m.group(2), int(m.group(3))
    month = int(month_tok) if month_tok.isdigit() else MONTHS.get(month_tok[:3].lower())
    try:
        dob = date(year, month or 0, day)
    except (ValueError, TypeError):
        return None
    age = (date.today() - dob).days / 365.25
    return dob.isoformat() if 15 <= age <= 80 else None


def _location(text: str, lines: List[str]) -> Optional[str]:
    m = re.search(r"(?:location|current location|address|city)\s*[:\-]\s*([^\n|,]+)", text, re.IGNORECASE)
    search_space = [m.group(1)] if m else []
    search_space += lines[:12]
    for chunk in search_space:
        for city in CITIES:
            if re.search(rf"\b{re.escape(city)}\b", chunk, re.IGNORECASE):
                return {"Bangalore": "Bengaluru", "Gurgaon": "Gurugram"}.get(city, city)
    return None


_INSTITUTION_RE = re.compile(r"universit|institut|college|\biit\b|\bnit\b|\biiit\b|\bbits\b|school of", re.IGNORECASE)


def _education(lines: List[str]) -> Dict[str, Any]:
    section = _section(lines, ["education", "academic", "academics", "qualifications"]) or lines
    best_rank, best = len(DEGREES), None
    for idx, line in enumerate(section):
        for rank, (pattern, label) in enumerate(DEGREES):
            if rank < best_rank and re.search(pattern, line, re.IGNORECASE):
                best_rank, best = rank, (label, idx)
                break
    if not best:
        return {}
    label, idx = best
    nearby = section[max(0, idx - 1): idx + 3]
    college = next((l for l in nearby if _INSTITUTION_RE.search(l)), None)
    years = [int(y) for l in nearby for y in re.findall(r"\b(19[8-9]\d|20[0-4]\d)\b", l)]
    result: Dict[str, Any] = {"highestQualification": label}
    if college:
        # Keep only the institution part of e.g. "B.Tech in CS, National Institute of Technology, Trichy, 2020"
        parts = [p.strip() for p in re.split(r"[,|]", college) if p.strip()]
        start = next((i for i, p in enumerate(parts) if _INSTITUTION_RE.search(p)), 0)
        kept = [p for p in parts[start:start + 2] if not re.fullmatch(r"(19|20)\d{2}", p)]
        result["educationCollege"] = ", ".join(kept)[:150]
    if years:
        result["graduationYear"] = max(years)
    return result


def _role(text: str, lines: List[str]) -> Dict[str, str]:
    result: Dict[str, str] = {}
    m = re.search(r"(?im)^\s*(?:current\s+)?designation\s*[:\-]\s*(.+)$", text)
    if m:
        result["currentDesignation"] = m.group(1).strip()[:100]
    m = re.search(r"(?im)^\s*(?:current\s+)?(?:company|organi[sz]ation|employer)\s*[:\-]\s*(.+)$", text)
    if m:
        result["currentCompany"] = m.group(1).strip()[:100]
    if len(result) == 2:
        return result

    m = re.search(r"(?:currently\s+)?working\s+as\s+(?:an?\s+)?([A-Za-z][\w /&.\-]{2,60}?)\s+(?:at|with|in)\s+([A-Z][\w .&\-]{1,60}?)(?:[.,;\n]|\s+since|\s+from|$)", text)
    if m:
        result.setdefault("currentDesignation", m.group(1).strip())
        result.setdefault("currentCompany", m.group(2).strip())
        return result

    # First role line in the experience section: "Title at Company" / "Title | Company" / "Company - Title"
    section = _section(lines, ["experience", "work experience", "professional experience", "employment history", "employment"])
    for line in section[:6]:
        line_wo_dates = _RANGE_RE.sub("", line).strip(" |,-–—")
        parts = re.split(r"\s+(?:at|@)\s+|\s*\|\s*|\s+[-–—]\s+|,\s+", line_wo_dates, maxsplit=1)
        if len(parts) != 2 or not all(parts):
            continue
        a, b = parts[0].strip(), parts[1].strip()
        a_is_title = any(w in a.lower() for w in TITLE_WORDS)
        b_is_title = any(w in b.lower() for w in TITLE_WORDS)
        if a_is_title and not b_is_title:
            result.setdefault("currentDesignation", a[:100]); result.setdefault("currentCompany", b[:100])
        elif b_is_title and not a_is_title:
            result.setdefault("currentDesignation", b[:100]); result.setdefault("currentCompany", a[:100])
        else:
            continue
        break
    return result


def _summary(lines: List[str]) -> Optional[str]:
    section = _section(lines, ["summary", "professional summary", "profile", "profile summary",
                               "objective", "career objective", "about", "about me"])
    text = " ".join(section).strip()
    return text[:1000] if len(text) >= 30 else None


SKILL_ALIASES = {"React.js": "React", "Golang": "Go", "HTML5": "HTML", "CSS3": "CSS", "Google Cloud": "GCP",
                 "Express.js": "Node.js"}


def _skills(text: str, lines: List[str]) -> List[str]:
    lower = text.lower()
    spans: List[Tuple[int, int, str]] = []
    for skill in SKILLS:
        if len(skill) <= 2:  # 'C', 'R', 'Go': only as standalone, case-sensitive tokens
            matches = re.finditer(rf"(?<![\w+#.]){re.escape(skill)}(?![\w+#])", text)
        else:
            matches = re.finditer(rf"(?<![\w+#.]){re.escape(skill.lower())}(?![\w+#])", lower)
        spans += [(m.start(), m.end(), skill) for m in matches]
    # Drop matches that are part of a longer matched skill ("CSS" inside "Tailwind CSS")
    standalone = [
        (s, e, skill) for s, e, skill in spans
        if not any(os <= s and e <= oe and (oe - os) > (e - s) for os, oe, _ in spans)
    ]
    skills: List[str] = []
    for _, _, skill in sorted(standalone):
        canonical = SKILL_ALIASES.get(skill, skill)
        if canonical.lower() not in (s.lower() for s in skills):
            skills.append(canonical)
    # Items written in the skills section that aren't in the vocabulary
    for line in _section(lines, ["skills", "technical skills", "key skills", "core competencies"]):
        line = re.sub(r"^[^:]{1,30}:\s*", "", line)
        for item in re.split(r"[,|•;·]", line):
            item = item.strip(" -*\t")
            if 2 <= len(item) <= 30 and len(item.split()) <= 4 and not re.search(r"\d{4}", item):
                item = SKILL_ALIASES.get(item, item)
                if item.lower() not in (s.lower() for s in skills):
                    skills.append(item)
    return skills[:25]


# ----------------------------------------------------------------------------- entry point
def parse_resume(text: str) -> Dict[str, Any]:
    """Returns only the fields that were found."""
    lines = _lines(text)
    fields: Dict[str, Any] = {}

    first, last = _name(lines)
    if first:
        fields["firstName"], fields["lastName"] = first, last or ""

    m = re.search(r"[\w.+\-]+@[\w\-]+\.[\w.\-]+", text)
    if m:
        fields["email"] = m.group(0).lower()
    m = re.search(r"(?:\+?91[\s\-]?)?([6-9]\d{4}[\s\-]?\d{5})\b", text)
    if m:
        fields["phone"] = re.sub(r"\D", "", m.group(1))

    years, months = _experience(text, lines)
    if years is not None:
        fields["totalExperienceYears"], fields["totalExperienceMonths"] = years, months or 0

    fields.update(_role(text, lines))

    current, expected = _ctc(text, r"(?:current|present)"), _ctc(text, r"expected")
    if current is not None:
        fields["currentCtcLpa"] = current
    if expected is not None:
        fields["expectedCtcLpa"] = expected

    for key, value in (("noticePeriod", _notice(text)), ("dateOfBirth", _dob(text)),
                       ("location", _location(text, lines)), ("summary", _summary(lines))):
        if value:
            fields[key] = value

    fields.update(_education(lines))

    for key, pattern in (("linkedinUrl", r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w\-%]+/?"),
                         ("githubUrl", r"(?:https?://)?(?:www\.)?github\.com/[\w\-]+/?")):
        m = re.search(pattern, text, re.IGNORECASE)
        if m:
            url = m.group(0)
            fields[key] = url if url.startswith("http") else f"https://{url}"

    skills = _skills(text, lines)
    if skills:
        fields["skills"] = skills
    return fields
